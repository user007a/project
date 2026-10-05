import docx, re, openpyxl
from collections import Counter

SRC = "长期照护师五级-JX.docx"
OUT = "长期照护师五级-JX-导入.xlsx"

HEADERS = ['试题分类', '试题类型', '试题难度', '试题内容', '答案',
           '解析', '分值', '选项1', '选项2', '选项3', '选项4',
           '选项5', '选项6', '选项7', '选项8']

CAT = '长期照护师五级-JX'
DIFF = '中等'
SCORE = 2  # 要求10：分值都填2
# 要求9：解析不填 -> 固定 None

header_pat = re.compile(r'^\s*(\d+)\.\s*(.*?)\s*【(.+?)】\s*$')
opt_pat = re.compile(r'^([A-Ha-h])[\.、]\s*(.*)$')
ans_pat = re.compile(r'参考答案[：:]\s*(.+)$')

# 判断题：A=对/正确 -> "1"，B=错/错误 -> "0"
JUDGE_MAP = {'A': '1', 'B': '0'}

d = docx.Document(SRC)
lines = [p.text.strip() for p in d.paragraphs]

questions = []
current = None
stray = []  # 无法归类的非空行

for line in lines:
    if not line:
        continue
    hm = header_pat.match(line)
    if hm:
        if current:
            questions.append(current)
        current = {
            'num': hm.group(1),
            'content': hm.group(2).strip(),
            'type': hm.group(3).strip(),
            'options': [],
            'answer': None,
        }
        continue
    if current is None:
        stray.append(line)
        continue
    am = ans_pat.search(line)
    if am:
        val = am.group(1).strip()
        # 取首位字母作为答案（兼容 "A  (答案根据医学常识推断...)" 这类后缀说明）
        m = re.match(r'^([A-Ha-h])', val)
        current['answer'] = m.group(1).upper() if m else val
        continue
    om = opt_pat.match(line)
    if om:
        current['options'].append((om.group(1).upper(), om.group(2).strip()))
        continue
    # 既不是选项也不是答案：视为题干续行（合并进 content）
    current['content'] += '\n' + line

if current:
    questions.append(current)

# ---- 生成输出 ----
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Sheet1"
ws.append(HEADERS)

errors = []
type_cnt = Counter()
for q in questions:
    t = q['type']
    type_cnt[t] += 1
    content = q['content']
    opts = q['options']
    ans_raw = q['answer']
    out_opts = [None] * 8

    if t == '单选题':
        # 答案字母 -> 位置数字 1..8
        if ans_raw is None or not re.match(r'^[A-H]$', ans_raw):
            errors.append((q['num'], ans_raw, '单选题答案异常'))
            ans = ans_raw
        else:
            ans = str(ord(ans_raw) - ord('A') + 1)
        # 选项搬入 选项1..8
        for idx, (_, txt) in enumerate(opts[:8]):
            out_opts[idx] = txt
    elif t == '判断题':
        if ans_raw in JUDGE_MAP:
            ans = JUDGE_MAP[ans_raw]
        else:
            errors.append((q['num'], ans_raw, '判断题答案非A/B'))
            ans = ans_raw
        out_opts = [None] * 8  # 判断题无需选项
    else:
        errors.append((q['num'], t, '不支持的题型'))
        ans = ans_raw

    row = [CAT, t, DIFF, content, ans, None, SCORE] + out_opts
    ws.append(row)

wb.save(OUT)

print(f"输出文件: {OUT}")
print(f"解析题数: {len(questions)}  (单选题 {type_cnt.get('单选题',0)} / 判断题 {type_cnt.get('判断题',0)})")
print(f"已写入数据行: {ws.max_row - 1}")
print(f"异常行数: {len(errors)}")
for e in errors[:30]:
    print("  异常:", e)
if stray:
    print(f"⚠️ 未归类行数: {len(stray)}，示例:", stray[:5])
