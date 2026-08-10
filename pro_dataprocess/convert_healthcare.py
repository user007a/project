# -*- coding: utf-8 -*-
"""将 健康照护师（长期照护师）五级初级工.xlsx 转换为导入格式。

输出列：试题分类 | 试题类型 | 试题难度 | 试题内容 | 答案 | 解析 | 分值 | 选项1-8
参考：老年人能力评估师三级题库-导入.xlsx
"""
import openpyxl
from openpyxl.styles import Font, Alignment
from copy import copy

SRC = r'd:\dev\GitHub\project\pro_dataprocess\健康照护师（长期照护师）五级初级工.xlsx'
DST = r'd:\dev\GitHub\project\pro_dataprocess\健康照护师（长期照护师）五级初级工_导入.xlsx'

CATEGORY = '健康照护师（长期照护师）五级初级工'
DIFFICULTY = '中等'
SCORE = 5

HEADERS = ['试题分类', '试题类型', '试题难度', '试题内容', '答案', '解析', '分值',
           '选项1', '选项2', '选项3', '选项4', '选项5', '选项6', '选项7', '选项8']

LETTER_TO_NUM = {'A': 1, 'B': 2, 'C': 3, 'D': 4, 'E': 5, 'F': 6, 'G': 7, 'H': 8}


def norm(v):
    if v is None:
        return ''
    return str(v).strip()


def convert_answer(qtype, raw):
    """按题型转换答案。返回 (答案字符串, 是否异常)。"""
    s = norm(raw).upper().replace('，', ',').replace(' ', '')
    if qtype == '单选题':
        # 单个字母 -> 数字
        if s in LETTER_TO_NUM:
            return str(LETTER_TO_NUM[s]), False
        # 兼容已经是数字的情况
        if s.isdigit() and 1 <= int(s) <= 8:
            return s, False
        return s, True
    if qtype == '多选题':
        # 多个字母 -> 逗号分隔数字，按字母顺序排序
        letters = [ch for ch in s if ch in LETTER_TO_NUM]
        if not letters:
            return s, True
        letters = sorted(set(letters), key=lambda c: LETTER_TO_NUM[c])
        return ','.join(str(LETTER_TO_NUM[c]) for c in letters), False
    if qtype == '判断题':
        # A=对 -> 1, B=错 -> 0；兼容 已是 1/0、对/错、正确/错误、√/×、T/F
        if s in ('A', '对', '正确', '√', 'T', 'TRUE', 'Y', 'YES', '1'):
            return '1', False
        if s in ('B', '错', '错误', '×', 'X', 'F', 'FALSE', 'N', 'NO', '0'):
            return '0', False
        return s, True
    return s, True


def main():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    ws = wb['Sheet1']

    out_wb = openpyxl.Workbook()
    out_ws = out_wb.active
    out_ws.title = 'Sheet1'

    # 表头
    out_ws.append(HEADERS)
    for c in range(1, len(HEADERS) + 1):
        cell = out_ws.cell(1, c)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    anomalies = []
    type_counts = {}
    written = 0

    for r in range(2, ws.max_row + 1):
        qtype = norm(ws.cell(r, 6).value)
        stem = norm(ws.cell(r, 7).value)
        # 跳过空行
        if not qtype and not stem:
            continue
        # 源文件选项列：Col8=A, Col9=B, Col10=C, Col11=D, Col12=E（最多5个选项）
        # Col13=参考答案, Col14=解析（不是选项）
        opts = [norm(ws.cell(r, 8 + i).value) for i in range(5)] + [''] * 3
        raw_ans = ws.cell(r, 13).value
        analysis = norm(ws.cell(r, 14).value)

        ans, bad = convert_answer(qtype, raw_ans)
        if bad:
            anomalies.append((r, qtype, raw_ans, '答案格式异常'))

        type_counts[qtype] = type_counts.get(qtype, 0) + 1

        # 选项：判断题不填选项（与参考文件一致）
        if qtype == '判断题':
            opt_row = ['' for _ in range(8)]
        else:
            opt_row = opts[:8]  # 已含5个真实选项 + 3个空

        row = [CATEGORY, qtype, DIFFICULTY, stem, ans, analysis, SCORE] + opt_row
        out_ws.append(row)
        written += 1

    # 列宽
    widths = [28, 10, 10, 50, 12, 50, 8, 22, 22, 22, 22, 22, 22, 22, 22]
    for i, w in enumerate(widths, 1):
        out_ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    out_ws.row_dimensions[1].height = 24

    out_wb.save(DST)

    print(f'输出文件: {DST}')
    print(f'写入数据行: {written}')
    print(f'题型分布: {type_counts}')
    if anomalies:
        print(f'异常 {len(anomalies)} 条:')
        for a in anomalies[:20]:
            print(f'  源行{a[0]} | {a[1]} | 答案={a[2]} | {a[3]}')
    else:
        print('无异常')


if __name__ == '__main__':
    main()
