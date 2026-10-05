import openpyxl

SRC = "老年人能力评估师三级-2.xlsx"
OUT = "老年人能力评估师三级-2-导入.xlsx"

# --- 模板列头（与现有导入文件一致）---
HEADERS = ['试题分类', '试题类型', '试题难度', '试题内容', '答案',
           '解析', '分值', '选项1', '选项2', '选项3', '选项4',
           '选项5', '选项6', '选项7', '选项8']

CAT = '老年人能力评估师三级'
DIFF = '中等'
SCORE = 5

# 判断题 选项文本映射：A=对(正确)->1, B=错(错误)->0
JUDGE_MAP = {'A': '1', 'B': '0'}

def letter_to_pos(letter):
    """将选项字母 A-H 转为 1-8 的位置数字。"""
    s = str(letter).strip().upper()
    if len(s) == 1 and 'A' <= s <= 'H':
        return str(ord(s) - ord('A') + 1)
    return s  # 兜底：非预期内容原样返回

wb = openpyxl.load_workbook(SRC, data_only=True)
ws = wb["Sheet1"]

out_wb = openpyxl.Workbook()
out_ws = out_wb.active
out_ws.title = "Sheet1"
out_ws.append(HEADERS)

skipped = 0
errors = []
for r in range(2, ws.max_row + 1):
    qtype = ws.cell(row=r, column=3).value
    if qtype is None:
        continue
    qtype = str(qtype).strip()
    content = ws.cell(row=r, column=4).value
    raw_ans = ws.cell(row=r, column=5).value
    explain = ws.cell(row=r, column=6).value
    # 源选项列 G..N -> 选项1..8
    src_opts = [ws.cell(row=r, column=c).value for c in range(7, 15)]

    if qtype == '单选题':
        ans = letter_to_pos(raw_ans)
        opts = src_opts
    elif qtype == '判断题':
        key = str(raw_ans).strip().upper()
        ans = JUDGE_MAP.get(key)
        if ans is None:
            errors.append((r, raw_ans, '判断题答案非A/B'))
            ans = raw_ans
        opts = [None] * 8  # 判断题无需选项
    else:
        # 题目要求仅支持 单/多/填/判/问答；本题库仅含 单/判
        errors.append((r, qtype, '不支持的题型'))
        skipped += 1
        continue

    row = [CAT, qtype, DIFF, content, ans, explain, SCORE] + list(opts)
    out_ws.append(row)

out_wb.save(OUT)

print(f"输出文件: {OUT}")
print(f"数据行数: {out_ws.max_row - 1}")
print(f"跳过行数: {skipped}")
print(f"异常行数: {len(errors)}")
for e in errors[:20]:
    print("  异常:", e)
